from datetime import datetime

import pytest

from hhmon import filters
from hhmon.exporters import to_csv, to_excel
from hhmon.filters import FilterConfig
from hhmon.hh_client import HHError, parse_state, parse_vacancies
from hhmon.notify import build_digest, format_vacancy
from hhmon.storage import Storage


# --- разбор страницы -------------------------------------------------------------

def test_parse_state_extracts_json(page_html):
    state = parse_state(page_html)
    assert state["vacancySearchResult"]["totalResults"] == 91


def test_parse_state_raises_without_template():
    with pytest.raises(HHError):
        parse_state("<html><body>ничего нет</body></html>")


def test_parse_vacancies_fields(vacancies):
    v = vacancies[0]
    assert v.id == 137196945
    assert v.company == "Topface Media"
    assert (v.salary_from, v.salary_to, v.currency) == (40000, 60000, "RUR")
    assert v.salary_text == "от 40 000 до 60 000 ₽ на руки"
    assert v.experience_ru == "без опыта"
    assert v.employment_ru == "полная"
    assert v.work_formats == ["REMOTE"]
    assert v.published_at == datetime.fromtimestamp(1789743222)
    assert v.responses == 215
    assert v.no_resume_needed and v.company_trusted and v.accredited_it


def test_no_compensation_rendered(vacancies):
    assert vacancies[2].salary_text == "не указана"


def test_parse_skips_broken_record(page_html):
    state = parse_state(page_html)
    state["vacancySearchResult"]["vacancies"].append({"name": "без id"})
    parsed, total = parse_vacancies(state)
    assert len(parsed) == 3 and total == 91


# --- фильтры ---------------------------------------------------------------------

def test_remote_only_drops_hybrid(vacancies):
    kept = filters.apply(vacancies, FilterConfig(remote_only=True))
    assert {v.id for v in kept} == {137196945, 136996674}


def test_exclude_keywords_and_experience(vacancies):
    cfg = FilterConfig(remote_only=False, exclude_keywords=["senior"], allowed_experience=["noExperience"])
    kept = filters.apply(vacancies, cfg)
    assert all(v.experience == "noExperience" for v in kept)
    assert not any("senior" in v.name.lower() for v in kept)


def test_min_salary_uses_best_known_number(vacancies):
    cfg = FilterConfig(remote_only=False, min_salary=50_000)
    kept = {v.id for v in filters.apply(vacancies, cfg)}
    assert 137196945 in kept          # to=60000 проходит
    assert 137492681 in kept          # зарплата не указана — не отсекаем
    assert 136996674 not in kept      # 900 $ < 50 000


def test_max_responses(vacancies):
    kept = filters.apply(vacancies, FilterConfig(remote_only=False, max_responses=500))
    assert 136996674 not in {v.id for v in kept}


def test_suspicious_flags(vacancies):
    cfg = FilterConfig(
        remote_only=False,
        suspicious_companies=["ГОРЕЛКИНА"],
        suspicious_title_patterns=[r"оператор (онлайн-)?чата"],
    )
    kept = {v.id: v for v in filters.apply(vacancies, cfg)}
    assert "оплата в $ без опыта" in kept[136996674].flags
    assert "компания из списка подозрительных" in kept[136996674].flags
    assert "шаблонное название" in kept[136996674].flags
    assert kept[137196945].flags == []


def test_dedupe_merges_queries(vacancies):
    a, b = vacancies[0], vacancies[0].__class__(**{**vacancies[0].__dict__, "query": "n8n"})
    merged = filters.dedupe([a, b, vacancies[1]])
    assert len(merged) == 2
    assert merged[0].query == "test, n8n"


# --- хранилище -------------------------------------------------------------------

def test_storage_detects_new_only_once(tmp_path, vacancies):
    with Storage(tmp_path / "db.sqlite3") as db:
        first = db.upsert(vacancies)
        second = db.upsert(vacancies)
        assert len(first) == 3 and second == []
        db.set_status(vacancies[0].id, "applied")
        db.save_cover_letter(vacancies[0].id, "Здравствуйте!")
        stats = db.stats()
        assert stats["total"] == 3
        assert stats["by_status"] == {"applied": 1, "new": 2}
        row = next(r for r in db.all_rows() if r["id"] == vacancies[0].id)
        assert row["cover_letter"] == "Здравствуйте!"


# --- экспорт и уведомления -------------------------------------------------------

def test_excel_and_csv_export(tmp_path, vacancies):
    xlsx = to_excel(vacancies, tmp_path / "out" / "v.xlsx", new_ids={vacancies[0].id})
    csv_path = to_csv(vacancies, tmp_path / "out" / "v.csv")
    assert xlsx.exists() and xlsx.stat().st_size > 5_000
    text = csv_path.read_text(encoding="utf-8-sig")
    assert "Topface Media" in text and text.count("\n") == 4

    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb["Вакансии"]
    assert ws["A1"].value == "Вакансия"
    assert ws.max_row == 4
    assert ws.cell(row=2, column=12).hyperlink is not None
    assert "Сводка" in wb.sheetnames


def test_digest_is_split_by_telegram_limit(vacancies):
    many = [vacancies[0].__class__(**{**vacancies[0].__dict__, "id": i}) for i in range(60)]
    messages = build_digest(many)
    assert len(messages) > 1
    assert all(len(m) <= 4096 for m in messages)
    assert "Topface Media" in format_vacancy(vacancies[0])
    assert build_digest([]) == []


def test_zero_total_drops_similar_fillers(page_html):
    state = parse_state(page_html)
    state["vacancySearchResult"]["totalResults"] = 0
    parsed, total = parse_vacancies(state, query="вайбкодер")
    assert parsed == [] and total == 0
