"""Экспорт в Excel и CSV.

Excel делается «под человека»: закреплённая шапка, автофильтр, кликабельные ссылки,
подсветка подозрительных строк, лист со сводкой по запросам и опыту.
"""

from __future__ import annotations

import csv
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .hh_client import Vacancy

COLUMNS = [
    ("Вакансия", "name", 48),
    ("Компания", "company", 30),
    ("Зарплата", "salary", 26),
    ("Опыт", "experience", 12),
    ("Занятость", "employment", 12),
    ("Формат", "work_formats", 12),
    ("Регион", "area", 18),
    ("Опубликовано", "published_at", 17),
    ("Откликов", "responses", 10),
    ("Запрос", "query", 22),
    ("Пометки", "flags", 36),
    ("Ссылка", "url", 34),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
NEW_FILL = PatternFill("solid", fgColor="E2EFDA")
LINK_FONT = Font(color="0563C1", underline="single")


def to_excel(vacancies: list[Vacancy], path: str | Path, new_ids: set[int] | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    new_ids = new_ids or set()

    wb = Workbook()
    ws = wb.active
    ws.title = "Вакансии"

    for col, (title, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    ordered = sorted(vacancies, key=lambda v: v.published_at, reverse=True)
    for row_idx, v in enumerate(ordered, start=2):
        row = v.to_row()
        for col, (_, key, _) in enumerate(COLUMNS, start=1):
            value = row[key]
            cell = ws.cell(row=row_idx, column=col, value=value)
            if key == "published_at":
                cell.number_format = "DD.MM.YYYY HH:MM"
            elif key == "url":
                cell.hyperlink = value
                cell.font = LINK_FONT
            if v.flags:
                cell.fill = WARN_FILL
            elif v.id in new_ids:
                cell.fill = NEW_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(ordered) + 1, 2)}"

    summary = wb.create_sheet("Сводка")
    summary["A1"], summary["B1"] = "Всего вакансий", len(vacancies)
    summary["A2"], summary["B2"] = "Новых за этот запуск", len([v for v in vacancies if v.id in new_ids])
    summary["A3"], summary["B3"] = "С пометками", len([v for v in vacancies if v.flags])
    summary["A4"], summary["B4"] = "Сформировано", datetime.now().strftime("%d.%m.%Y %H:%M")

    r = 6
    for title, counter in (
        ("По запросу", Counter(q.strip() for v in vacancies for q in v.query.split(","))),
        ("По опыту", Counter(v.experience_ru for v in vacancies)),
        ("По занятости", Counter(v.employment_ru for v in vacancies)),
    ):
        summary.cell(row=r, column=1, value=title).font = Font(bold=True)
        r += 1
        for key, count in counter.most_common():
            summary.cell(row=r, column=1, value=key)
            summary.cell(row=r, column=2, value=count)
            r += 1
        r += 1
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 12

    wb.save(path)
    return path


def to_csv(vacancies: list[Vacancy], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=[key for _, key, _ in COLUMNS])
        writer.writeheader()
        for v in sorted(vacancies, key=lambda v: v.published_at, reverse=True):
            row = v.to_row()
            row["published_at"] = row["published_at"].strftime("%Y-%m-%d %H:%M")
            writer.writerow({key: row[key] for _, key, _ in COLUMNS})
    return path
